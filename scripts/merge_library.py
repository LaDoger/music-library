#!/usr/bin/env python3
"""Merge parts/*_rows.csv into library.csv + library.xlsx.

- concatenates codex, claude, grok part files (exact SCHEMA header)
- drops exact/near duplicate pieces (logged to parts/MERGE_NOTES.md by caller)
- renames doubled filenames (<id>_<id>.ext, <id>_<composer>_<short>.ext) to
  <id>.ext (keeping _lilypond/_midi tails) on disk and in every CSV cell
- re-renders previews that are missing or more than 1 LU off -16 LUFS
- writes library.csv and library.xlsx

Run from anywhere:  python3 scripts/merge_library.py
Safe to re-run: already-clean names are left alone.
"""
from __future__ import annotations

import csv
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from make_preview import integrated, make_preview, parse_time  # noqa: E402

PARTS = ["codex", "claude", "grok"]
PATH_COLS = ["local_score_path", "local_audio_path", "preview_path"]
KEEP_TAILS = ("_lilypond", "_midi")


def header() -> list[str]:
    for line in (ROOT / "SCHEMA.md").read_text(encoding="utf-8").splitlines():
        if line.startswith("id,composer,"):
            return line.strip().split(",")
    raise SystemExit("SCHEMA header not found")


def load(cols):
    rows = []
    for part in PARTS:
        path = ROOT / "parts" / f"{part}_rows.csv"
        with path.open(newline="", encoding="utf-8") as fh:
            reader = csv.DictReader(fh)
            assert reader.fieldnames == cols, f"{path} header mismatch"
            rows += list(reader)
    return rows


def dedupe(rows):
    """Drop rows sharing an id or composer+catalog+movement; keep the better one."""
    def norm(s):
        return re.sub(r"[^a-z0-9]+", "", (s or "").lower())

    def score(r):
        filled = sum(bool(r[c].strip()) for c in r)
        return (r["verified"] == "yes", bool(r["local_audio_path"]), bool(r["local_score_path"]), filled)

    kept, drops = {}, []
    for r in rows:
        keys = {("id", r["id"])}
        if norm(r["catalog"]) and norm(r["movement"]):
            keys.add(("cat", norm(r["composer"]), norm(r["catalog"]), norm(r["movement"])))
        hit = next((k for k in keys if k in kept), None)
        if hit is None:
            for k in keys:
                kept[k] = r
            continue
        other = kept[hit]
        winner, loser = (r, other) if score(r) > score(other) else (other, r)
        for k in list(kept):
            if kept[k] is other:
                kept[k] = winner
        drops.append((loser["id"], loser["added_by"], winner["id"], winner["added_by"]))
    seen, out = set(), []
    for r in rows:
        if any(v is r for v in kept.values()) and id(r) not in seen:
            seen.add(id(r))
            out.append(r)
    return out, drops


def clean_name(rel: str, rid: str) -> str:
    p = Path(rel)
    stem, suffix = p.stem, p.suffix
    if not stem.startswith(rid + "_"):
        return rel
    rest = stem[len(rid):]
    first = rid.split("_")[0]
    if not rest.startswith("_" + first):
        return rel  # not a doubled name
    tail = next((t for t in KEEP_TAILS if rest.endswith(t)), "")
    return str(p.with_name(rid + tail + suffix))


def rename_all(rows):
    mapping = {}
    for r in rows:
        for col in PATH_COLS:
            for rel in filter(None, (s.strip() for s in r[col].split(";"))):
                new = clean_name(rel, r["id"])
                if new != rel:
                    mapping[rel] = new
    # sibling originals on disk that no row references (codex .ly / zips / m4a)
    ids = sorted({r["id"] for r in rows}, key=len, reverse=True)
    for f in sorted((ROOT / "files").glob("*/*")) + sorted((ROOT / "previews").glob("*.mp3")):
        rel = str(f.relative_to(ROOT))
        if rel in mapping:
            continue
        rid = next((i for i in ids if f.name.startswith(i + "_")), None)
        if rid:
            new = clean_name(rel, rid)
            if new != rel:
                mapping[rel] = new
    for old, new in sorted(mapping.items()):
        src, dst = ROOT / old, ROOT / new
        if src.exists():
            if dst.exists():
                raise SystemExit(f"refusing to overwrite {new}")
            src.rename(dst)
    # rewrite every cell (paths and any mention of old basenames)
    repl = sorted(mapping.items(), key=lambda kv: len(kv[0]), reverse=True)
    for r in rows:
        for c in r:
            v = r[c]
            for old, new in repl:
                if old in v:
                    v = v.replace(old, new)
            for old, new in repl:
                ob, nb = Path(old).name, Path(new).name
                if ob in v:
                    v = v.replace(ob, nb)
            r[c] = v
    return mapping


def preview_source(r):
    audio = [s for s in r["local_audio_path"].split(";") if s.strip()]
    scores = [s for s in r["local_score_path"].split(";") if s.strip().endswith(".mid")]
    m = re.search(r"(\d+:\d{2}(?::\d{2})?)", r["notable_excerpt"])
    start = parse_time(m.group(1)) if m else 0.0
    if audio:
        return ROOT / audio[0].strip(), False, start
    if scores:
        return ROOT / scores[0].strip(), True, start
    return None, False, 0.0


def fix_previews(rows):
    log = []
    for r in rows:
        if not (r["local_score_path"].strip() or r["local_audio_path"].strip()):
            continue
        want = r["preview_path"].strip() or f"previews/{r['id']}.mp3"
        r["preview_path"] = want
        out = ROOT / want
        reason = None
        if not out.is_file():
            reason = "missing"
        else:
            lufs = integrated(out)
            if abs(lufs + 16.0) > 1.0:
                reason = f"loudness {lufs:.1f} LUFS"
        if not reason:
            continue
        src, midi, start = preview_source(r)
        if src is None:
            log.append((r["id"], reason, "NO SOURCE (only zipped scores)"))
            continue
        got = make_preview(src, out, start, midi=midi)
        log.append((r["id"], reason, f"re-rendered from {src.relative_to(ROOT)} @ {start:.0f}s -> {got:.1f} LUFS"))
    return log


def write_xlsx(rows, cols, path):
    import pandas as pd
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    df = pd.DataFrame(rows, columns=cols)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="library")
        ws = xw.sheets["library"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(cols, 1):
            width = max(len(c), *(min(len(str(v)), 60) for v in df[c])) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 62)
        for cell in ws[1]:
            cell.font = Font(bold=True)


def main():
    cols = header()
    rows = load(cols)
    n_in = len(rows)
    rows, drops = dedupe(rows)
    mapping = rename_all(rows)
    plog = fix_previews(rows)
    rows.sort(key=lambda r: r["id"])
    with (ROOT / "library.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)
    write_xlsx(rows, cols, ROOT / "library.xlsx")
    map_file = ROOT / "logs" / "merge_rename_map.tsv"
    lines = set(map_file.read_text(encoding="utf-8").splitlines()) if map_file.exists() else set()
    lines |= {f"{old}\t{new}" for old, new in mapping.items()}
    map_file.write_text("".join(f"{l}\n" for l in sorted(lines)), encoding="utf-8")
    scores = sum(bool(r["local_score_path"].strip()) for r in rows)
    audio = sum(bool(r["local_audio_path"].strip()) for r in rows)
    prev = sum(bool(r["preview_path"].strip()) and (ROOT / r["preview_path"]).is_file() for r in rows)
    print(f"IN {n_in} OUT {len(rows)} DROPS {drops}")
    print(f"RENAMED {len(mapping)}")
    for line in plog:
        print("PREVIEW", *line, sep="\t")
    print(f"COUNTS N={len(rows)} S={scores} A={audio} P={prev}")


if __name__ == "__main__":
    main()
